"""Exports the exchange and VAT rates applied to a calculation."""

from __future__ import annotations

import io
from collections import defaultdict
from datetime import date
from decimal import Decimal
from typing import Iterable

from .ecb_rates import fixed_eur_rate, get_closing_rate, get_ioss_rate_date, get_oss_rate_date
from .i18n import _
from .models import Scenario, VatResult


def _transaction_date(value: str) -> date | None:
    try:
        return date.fromisoformat((value or "")[:10])
    except ValueError:
        return None


def _decimal(value: Decimal | None) -> str:
    return str(value) if value is not None else ""


def extract_rates_evidence_records(
    results: Iterable[VatResult],
    period: str,
    translator=None,
) -> tuple[
    list[tuple[str, Decimal, str, str]],
    list[tuple[str, str, str, str, str, Decimal | None, str]],
]:
    """Extract structured VAT and FX rate evidence records from calculation results.

    Returns:
        vat_records: list of (country, vat_rate_decimal, first_date_str, last_date_str)
        fx_records: list of (type_translated, regime, date_str, closing_date_str, currency, rate_decimal_or_none, source)
    """
    translate = translator or _
    daily_rates: set[tuple[str, str, str, str]] = set()
    vat_periods: dict[tuple[str, str], list[date]] = defaultdict(list)
    closing_rates: dict[tuple[str, str, date], list[VatResult]] = defaultdict(list)

    for result in results:
        sale = result.sale
        currency = (sale.original_currency or "EUR").upper()
        tx_date = _transaction_date(sale.transaction_date)
        tx_date_text = tx_date.isoformat() if tx_date else sale.transaction_date[:10]

        if currency != "EUR":
            daily_rates.add((
                tx_date_text,
                currency,
                _decimal(sale.exchange_rate),
                sale.exchange_rate_source,
            ))

        if result.vat_rate != 0:
            vat_key = (result.vat_country, str(result.vat_rate))
            if tx_date:
                vat_periods[vat_key].append(tx_date)
            else:
                vat_periods[vat_key]

        if currency == "EUR" or result.scenario not in (Scenario.OSS_B2C, Scenario.IOSS_DIRECT):
            continue
        effective_date = tx_date or date.today()
        regime = "IOSS" if result.scenario == Scenario.IOSS_DIRECT else "OSS"
        rate_date_fn = get_ioss_rate_date if regime == "IOSS" else get_oss_rate_date
        closing_date = rate_date_fn(period, effective_date)
        closing_rates[(regime, currency, closing_date)].append(result)

    fx_records: list[tuple[str, str, str, str, str, Decimal | None, str]] = []
    for tx_date_text, currency, rate_str, source in sorted(daily_rates):
        rate_dec = Decimal(rate_str) if rate_str else None
        fx_records.append((
            translate("rates_evidence_daily"),
            "",
            tx_date_text,
            "",
            currency,
            rate_dec,
            source,
        ))

    for (regime, currency, closing_date), rate_results in sorted(closing_rates.items()):
        closing_rate = fixed_eur_rate(currency, closing_date)
        if closing_rate is not None:
            source = "fixed_eu"
        else:
            closing_rate = get_closing_rate(currency, closing_date)
            source = "ecb_closing" if closing_rate is not None else ""

        if closing_rate is not None:
            fx_records.append((
                translate("rates_evidence_closing"),
                regime,
                "",
                closing_date.isoformat(),
                currency,
                closing_rate,
                source,
            ))
            continue

        seen_fallbacks: set[tuple[str, str]] = set()
        for result in rate_results:
            sale = result.sale
            tx_date_text = (sale.transaction_date or "")[:10]
            fallback_key = (tx_date_text, str(sale.exchange_rate))
            if fallback_key in seen_fallbacks:
                continue
            seen_fallbacks.add(fallback_key)
            fx_records.append((
                translate("rates_evidence_fallback"),
                regime,
                tx_date_text,
                closing_date.isoformat(),
                currency,
                sale.exchange_rate,
                sale.exchange_rate_source,
            ))

    vat_records: list[tuple[str, Decimal, str, str]] = []
    for (country, rate_str), dates in sorted(vat_periods.items()):
        first_d = min(dates).isoformat() if dates else ""
        last_d = max(dates).isoformat() if dates else ""
        vat_records.append((
            country,
            Decimal(rate_str),
            first_d,
            last_d,
        ))

    return vat_records, fx_records


def build_rates_evidence_xlsx(
    results: Iterable[VatResult],
    period: str,
    company_name: str = "",
    siren: str = "",
    translator=None,
) -> bytes:
    """Build an Excel workbook (.xlsx) with 2 tabs: 'TVA' and 'BCE'."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    translate = translator or _
    vat_records, fx_records = extract_rates_evidence_records(results, period, translator=translate)

    wb = Workbook()

    navy_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    zebra_fill = PatternFill(start_color="F2F5F8", end_color="F2F5F8", fill_type="solid")

    title_font = Font(name="Calibri", size=13, bold=True, color="1F4E79")
    meta_font = Font(name="Calibri", size=9, italic=True, color="595959")
    header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)

    thin_side = Side(style="thin", color="D9D9D9")
    cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    col_label = lambda key: translate(f"rates_evidence_col_{key}")

    meta_parts = []
    if company_name:
        meta_parts.append(f"{translate('vies_certificate_company')} : {company_name}")
    if siren:
        meta_parts.append(f"{translate('vies_certificate_siren')} : {siren}")
    meta_parts.append(f"{translate('vies_certificate_period')} : {period}")
    meta_str = " | ".join(meta_parts)

    # =========================================================
    # TAB 1: TVA
    # =========================================================
    ws_tva = wb.active
    ws_tva.title = "TVA"
    ws_tva.views.sheetView[0].showGridLines = True

    ws_tva.append([translate("rates_evidence_pdf_vat_header")])
    ws_tva.cell(row=1, column=1).font = title_font
    ws_tva.append([meta_str])
    ws_tva.cell(row=2, column=1).font = meta_font
    ws_tva.append([])

    tva_headers = [
        col_label("country"),
        col_label("vat_rate"),
        col_label("first_date"),
        col_label("last_date"),
    ]
    ws_tva.append(tva_headers)
    ws_tva.row_dimensions[4].height = 24
    for c_idx in range(1, len(tva_headers) + 1):
        cell = ws_tva.cell(row=4, column=c_idx)
        cell.font = header_font
        cell.fill = navy_fill
        cell.alignment = header_align
        cell.border = cell_border

    row_idx = 5
    for country, vat_rate_dec, first_date, last_date in vat_records:
        rate_val = float(vat_rate_dec) if vat_rate_dec is not None else None

        c_country = ws_tva.cell(row=row_idx, column=1, value=country)
        c_rate = ws_tva.cell(row=row_idx, column=2, value=rate_val)
        c_first = ws_tva.cell(row=row_idx, column=3, value=first_date)
        c_last = ws_tva.cell(row=row_idx, column=4, value=last_date)

        c_country.alignment = align_center
        c_rate.alignment = align_right
        c_rate.number_format = '0.0"%"'
        c_first.alignment = align_center
        c_last.alignment = align_center

        row_fill = zebra_fill if row_idx % 2 == 1 else None
        for cell in (c_country, c_rate, c_first, c_last):
            cell.font = data_font
            cell.border = cell_border
            if row_fill:
                cell.fill = row_fill
        row_idx += 1

    for c_idx in range(1, len(tva_headers) + 1):
        col_letter = get_column_letter(c_idx)
        max_len = 0
        for r in range(4, ws_tva.max_row + 1):
            val = ws_tva.cell(row=r, column=c_idx).value
            if val is not None:
                max_len = max(max_len, len(str(val)))
        ws_tva.column_dimensions[col_letter].width = max(max_len + 5, 18)

    # =========================================================
    # TAB 2: BCE
    # =========================================================
    ws_bce = wb.create_sheet(title="BCE")
    ws_bce.views.sheetView[0].showGridLines = True

    ws_bce.append([translate("rates_evidence_pdf_exchange_header")])
    ws_bce.cell(row=1, column=1).font = title_font
    ws_bce.append([meta_str])
    ws_bce.cell(row=2, column=1).font = meta_font
    ws_bce.append([])

    bce_headers = [
        col_label("type"),
        col_label("regime"),
        col_label("date"),
        col_label("closing_date"),
        col_label("currency"),
        col_label("rate"),
        col_label("source"),
    ]
    ws_bce.append(bce_headers)
    ws_bce.row_dimensions[4].height = 24
    for c_idx in range(1, len(bce_headers) + 1):
        cell = ws_bce.cell(row=4, column=c_idx)
        cell.font = header_font
        cell.fill = navy_fill
        cell.alignment = header_align
        cell.border = cell_border

    row_idx = 5
    for rate_type, regime, dt, closing_dt, currency, fx_rate_dec, source in fx_records:
        rate_val = float(fx_rate_dec) if fx_rate_dec is not None else None

        c_type = ws_bce.cell(row=row_idx, column=1, value=rate_type)
        c_regime = ws_bce.cell(row=row_idx, column=2, value=regime)
        c_dt = ws_bce.cell(row=row_idx, column=3, value=dt)
        c_closing_dt = ws_bce.cell(row=row_idx, column=4, value=closing_dt)
        c_curr = ws_bce.cell(row=row_idx, column=5, value=currency)
        c_rate = ws_bce.cell(row=row_idx, column=6, value=rate_val)
        c_src = ws_bce.cell(row=row_idx, column=7, value=source)

        c_type.alignment = align_left
        c_regime.alignment = align_center
        c_dt.alignment = align_center
        c_closing_dt.alignment = align_center
        c_curr.alignment = align_center
        c_rate.alignment = align_right
        if rate_val is not None:
            c_rate.number_format = "#,##0.0000"
        c_src.alignment = align_left

        row_fill = zebra_fill if row_idx % 2 == 1 else None
        for cell in (c_type, c_regime, c_dt, c_closing_dt, c_curr, c_rate, c_src):
            cell.font = data_font
            cell.border = cell_border
            if row_fill:
                cell.fill = row_fill
        row_idx += 1

    for c_idx in range(1, len(bce_headers) + 1):
        col_letter = get_column_letter(c_idx)
        max_len = 0
        for r in range(4, ws_bce.max_row + 1):
            val = ws_bce.cell(row=r, column=c_idx).value
            if val is not None:
                max_len = max(max_len, len(str(val)))
        ws_bce.column_dimensions[col_letter].width = max(max_len + 5, 18)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
