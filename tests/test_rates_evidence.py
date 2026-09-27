from datetime import date
from decimal import Decimal
from io import BytesIO

import openpyxl

from tva_intracom.models import BuyerType, Channel, Collector, Sale, Scenario, VatResult
from tva_intracom.rates_evidence import extract_rates_evidence_records, build_rates_evidence_xlsx
from tva_intracom.rates_evidence_pdf import generate_rates_evidence_pdf


def _result(
    sale_id: str,
    tx_date: str,
    vat_rate: Decimal = Decimal("19"),
    category: str = "STANDARD",
    currency: str = "USD",
    scenario: Scenario = Scenario.OSS_B2C,
) -> VatResult:
    sale = Sale(
        sale_id=sale_id,
        amount_ht=Decimal("100.00"),
        buyer_type=BuyerType.B2C,
        stock_country="FR",
        buyer_country="DE",
        original_currency=currency,
        original_amount=Decimal("110.00"),
        exchange_rate=Decimal("1.10"),
        exchange_rate_source="ecb",
        transaction_date=tx_date,
        product_category=category,
    )
    return VatResult(
        sale=sale,
        scenario=scenario,
        vat_country="DE",
        vat_rate=vat_rate,
        vat_amount=Decimal("100") * vat_rate / Decimal("100"),
        collector=Collector.SELLER,
        channel=Channel.OSS,
        note="",
    )


def test_rates_evidence_extracts_daily_closing_and_vat_rates(monkeypatch):
    monkeypatch.setattr(
        "tva_intracom.rates_evidence.get_closing_rate",
        lambda currency, closing_date: Decimal("1.12"),
    )
    vat_records, fx_records = extract_rates_evidence_records(
        [_result("S1", "2026-04-01"), _result("S2", "2026-04-03")],
        "2026-Q2",
    )
    daily = [r for r in fx_records if r[0] == "Change quotidien"]
    closing = [r for r in fx_records if r[0] == "Clôture BCE"]

    assert len(daily) == 2
    assert {r[2] for r in daily} == {"2026-04-01", "2026-04-03"}
    assert len(closing) == 1
    assert closing[0][1:7] == ("OSS", "", "2026-06-30", "USD", Decimal("1.12"), "ecb_closing")
    assert len(vat_records) == 1
    assert vat_records[0] == ("DE", Decimal("19"), "2026-04-01", "2026-04-03")


def test_rates_evidence_lists_country_rate_once_across_categories_and_years():
    early = _result("S1", "2021-01-08")
    later = _result("S2", "2026-05-26", category="FOOD")
    changed_rate = _result("S3", "2026-05-29", vat_rate=Decimal("7"))
    vat_records, _ = extract_rates_evidence_records([early, later, changed_rate], "2026-Q2")

    assert len(vat_records) == 2
    assert vat_records == [
        ("DE", Decimal("19"), "2021-01-08", "2026-05-26"),
        ("DE", Decimal("7"), "2026-05-29", "2026-05-29"),
    ]


def test_rates_evidence_marks_sale_rate_fallback(monkeypatch):
    monkeypatch.setattr(
        "tva_intracom.rates_evidence.get_closing_rate",
        lambda currency, closing_date: None,
    )
    _, fx_records = extract_rates_evidence_records([_result("S1", "2026-04-01")], "2026-Q2")
    fallback = [r for r in fx_records if r[0] == "Repli taux de vente"]

    assert len(fallback) == 1
    assert fallback[0][1:7] == ("OSS", "2026-04-01", "2026-06-30", "USD", Decimal("1.10"), "ecb")


def test_rates_evidence_keeps_non_oss_currency_on_daily_rate_only(monkeypatch):
    closing_rate_calls = []

    def unexpected_closing_rate(currency, closing_date):
        closing_rate_calls.append((currency, closing_date))
        return Decimal("0.86")

    monkeypatch.setattr(
        "tva_intracom.rates_evidence.get_closing_rate",
        unexpected_closing_rate,
    )
    _, fx_records = extract_rates_evidence_records(
        [_result("S1", "2026-04-01", currency="GBP", scenario=Scenario.DOMESTIC)],
        "2026-Q2",
    )

    assert len(fx_records) == 1
    assert fx_records[0][0] == "Change quotidien"
    assert fx_records[0][2:6] == ("2026-04-01", "", "GBP", Decimal("1.10"))
    assert closing_rate_calls == []


def test_rates_evidence_uses_monthly_closing_date_for_ioss(monkeypatch):
    closing_rate_calls = []

    def record_closing_rate(currency, closing_date):
        closing_rate_calls.append((currency, closing_date))
        return Decimal("1.12")

    monkeypatch.setattr(
        "tva_intracom.rates_evidence.get_closing_rate",
        record_closing_rate,
    )
    _, fx_records = extract_rates_evidence_records(
        [
            _result(
                "S1",
                "2026-04-01",
                scenario=Scenario.IOSS_DIRECT,
            )
        ],
        "2026-04",
    )
    closing = [row for row in fx_records if row[0] == "Clôture BCE"]

    assert len(closing) == 1
    assert closing[0][1:7] == (
        "IOSS",
        "",
        "2026-04-30",
        "USD",
        Decimal("1.12"),
        "ecb_closing",
    )
    assert closing_rate_calls == [("USD", date(2026, 4, 30))]


def test_rates_evidence_excludes_zero_vat_rates(monkeypatch):
    monkeypatch.setattr(
        "tva_intracom.rates_evidence.get_closing_rate",
        lambda currency, closing_date: Decimal("1.12"),
    )
    vat_records, _ = extract_rates_evidence_records(
        [_result("S0", "2026-04-01", vat_rate=Decimal("0"))],
        "2026-Q2",
    )

    assert len(vat_records) == 0


def test_rates_evidence_pdf_has_pdf_signature_and_unique_reference():
    results = [_result("S1", "2026-04-01")]
    options = {
        "company_name": "Example",
        "siren": "123456789",
        "scope_id": "private-scope",
    }

    first = generate_rates_evidence_pdf(results, "2026-Q2", **options)
    second = generate_rates_evidence_pdf(results, "2026-Q2", **options)

    assert first.startswith(b"%PDF-")
    assert first.endswith(b"%%EOF\n")
    assert first != second


def test_rates_evidence_xlsx_generates_two_sheets_with_expected_headers_and_data(monkeypatch):
    monkeypatch.setattr(
        "tva_intracom.rates_evidence.get_closing_rate",
        lambda currency, closing_date: Decimal("1.12"),
    )
    xlsx_bytes = build_rates_evidence_xlsx(
        [_result("S1", "2026-04-01"), _result("S2", "2026-04-03")],
        "2026-Q2",
        company_name="TestCo",
        siren="123456789",
    )

    assert xlsx_bytes.startswith(b"PK\x03\x04")

    wb = openpyxl.load_workbook(BytesIO(xlsx_bytes))
    assert wb.sheetnames == ["TVA", "BCE"]

    ws_tva = wb["TVA"]
    assert "TestCo" in str(ws_tva.cell(row=2, column=1).value)
    tva_headers = [ws_tva.cell(row=4, column=col).value for col in range(1, 5)]
    assert tva_headers == [
        "Pays TVA",
        "Taux TVA (%)",
        "Première date d'opération",
        "Dernière date d'opération",
    ]
    assert ws_tva.cell(row=5, column=1).value == "DE"
    assert ws_tva.cell(row=5, column=2).value == 19.0
    assert ws_tva.cell(row=5, column=3).value == "2026-04-01"
    assert ws_tva.cell(row=5, column=4).value == "2026-04-03"

    ws_bce = wb["BCE"]
    assert "TestCo" in str(ws_bce.cell(row=2, column=1).value)
    bce_headers = [ws_bce.cell(row=4, column=col).value for col in range(1, 8)]
    assert bce_headers == [
        "Type",
        "Régime",
        "Date du taux / opération",
        "Clôture OSS/IOSS",
        "Devise",
        "Cours (1 EUR = devise) / taux TVA (%)",
        "Source",
    ]
    daily_rows = [row for row in ws_bce.iter_rows(min_row=5, values_only=True) if row[0] == "Change quotidien"]
    closing_rows = [row for row in ws_bce.iter_rows(min_row=5, values_only=True) if row[0] == "Clôture BCE"]

    assert len(daily_rows) == 2
    assert len(closing_rows) == 1
    assert closing_rows[0][1:7] == ("OSS", None, "2026-06-30", "USD", 1.12, "ecb_closing")
